"""Shared, non-test fixtures for the reporting suites.

Deliberately not named ``test_*``: several T13 modules import it, and importing
shared data from another test module would couple those modules together.

The small reconciliation result built here is the one the golden files describe.
It is deliberately tiny and hand-checkable — one group per outcome, one
unresolved row per category — so a golden diff points at a specific rule rather
than at three hundred rows of fixture noise.
"""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from las.reconcile import (
    REASON_INVALID_AMOUNT,
    REASON_INVALID_PERIOD,
    REASON_INVALID_UNIT,
    REASON_MATCH_INELIGIBLE,
    REASON_MATCH_NEEDS_REVIEW,
    REASON_MATCH_UNMATCHED,
    BatchReconcileResult,
    Contribution,
    GroupKey,
    GroupResult,
    ReconcileOutcome,
    Side,
    SideAggregate,
    UnresolvedRow,
    exact_sum,
    summarize,
)
from las.report import OutputFormat, ReportPolicy

GOLDEN_DIR = Path(__file__).resolve().parent / "golden" / "report"

#: The bands the sample configuration declares, as exact decimals.
SAMPLE_BANDS: tuple[Decimal, ...] = (Decimal(0), Decimal(1), Decimal(5))
RECIPIENTS: tuple[str, ...] = ("ops-team@example.com",)


def make_policy(
    formats: tuple[OutputFormat, ...] = (
        OutputFormat.XLSX,
        OutputFormat.JSON,
        OutputFormat.HTML,
    ),
    *,
    recipients: tuple[str, ...] = RECIPIENTS,
    variance_bands: tuple[Decimal, ...] = SAMPLE_BANDS,
    streaming_row_threshold: int = 50_000,
) -> ReportPolicy:
    """A policy mirroring the committed sample configuration."""
    return ReportPolicy(
        formats=formats,
        recipients=recipients,
        variance_bands=variance_bands,
        streaming_row_threshold=streaming_row_threshold,
    )


def contribution(
    side: Side = Side.FEED,
    amount: str = "10.00",
    unit: str = "EA",
    source: str = "alpha.csv",
    source_row: int = 2,
    row_index: int = 0,
) -> Contribution:
    return Contribution(
        side=side,
        amount=Decimal(amount),
        unit=unit,
        source=source,
        source_row=source_row,
        row_index=row_index,
    )


def aggregate(
    side: Side, *pairs: tuple[str, str], source_row: int = 2
) -> SideAggregate:
    """A side aggregate from ``(amount, unit)`` pairs; no pairs means absent."""
    contributions = tuple(
        contribution(side, amount, unit, source_row=source_row + offset, row_index=offset)
        for offset, (amount, unit) in enumerate(pairs)
    )
    by_unit: dict[str, list[Decimal]] = {}
    for item in contributions:
        by_unit.setdefault(item.unit, []).append(item.amount)
    # T12's own exact summation, not ``+``: the operator rounds under the
    # default decimal context, which would quietly shorten a long total in a
    # test built to prove long totals survive.
    totals = {unit: exact_sum(amounts) for unit, amounts in by_unit.items()}
    return SideAggregate(side=side, contributions=contributions, totals_by_unit=totals)


def group(
    account_ref: str,
    period: str,
    outcome: ReconcileOutcome,
    reason: str,
    feed: SideAggregate,
    system: SideAggregate,
    *,
    signed: str | None = None,
    absolute: str | None = None,
    band: str | None = None,
) -> GroupResult:
    return GroupResult(
        key=GroupKey(account_ref, period),
        feed=feed,
        system=system,
        outcome=outcome,
        reason=reason,
        signed_difference=None if signed is None else Decimal(signed),
        absolute_difference=None if absolute is None else Decimal(absolute),
        variance_band=None if band is None else Decimal(band),
    )


def unresolved(
    reason: str,
    source_row: int,
    *,
    source: str = "alpha.csv",
    worksheet: str | None = None,
    match_outcome: str | None = None,
    account_ref: str | None = None,
) -> UnresolvedRow:
    return UnresolvedRow(
        side=Side.FEED,
        reason=reason,
        source=source,
        source_row=source_row,
        row_index=source_row - 2,
        worksheet=worksheet,
        match_outcome=match_outcome,
        account_ref=account_ref,
    )


def golden_result() -> BatchReconcileResult:
    """One small reconciliation result covering every outcome and category.

    Six groups — one per outcome — and six unresolved rows, one per reason.
    Every number is short enough to check by eye, which is the point of a
    golden file.
    """
    groups = (
        group(
            "ACC-100000", "2026-01", ReconcileOutcome.BALANCED, "totals_equal",
            aggregate(Side.FEED, ("10.00", "EA")),
            aggregate(Side.SYSTEM, ("10.00", "EA")),
            signed="0", absolute="0", band="0",
        ),
        group(
            "ACC-100000", "2026-02", ReconcileOutcome.WITHIN_TOLERANCE,
            "within_variance_band",
            aggregate(Side.FEED, ("10.50", "EA")),
            aggregate(Side.SYSTEM, ("10.00", "EA")),
            signed="0.50", absolute="0.50", band="1",
        ),
        group(
            "ACC-100001", "2026-01", ReconcileOutcome.MISMATCH,
            "outside_variance_bands",
            aggregate(Side.FEED, ("99.25", "KG")),
            aggregate(Side.SYSTEM, ("10.00", "KG")),
            signed="89.25", absolute="89.25",
        ),
        group(
            "ACC-100001", "2026-02", ReconcileOutcome.UNIT_MISMATCH,
            "unit_differs_across_sides",
            aggregate(Side.FEED, ("3.00", "EA")),
            aggregate(Side.SYSTEM, ("4.00", "KG")),
        ),
        group(
            "ACC-100002", "2026-01", ReconcileOutcome.MISSING_IN_FEED,
            "no_feed_contribution",
            SideAggregate(side=Side.FEED),
            aggregate(Side.SYSTEM, ("7.00", "BOX")),
        ),
        group(
            "ACC-100003", "2026-01", ReconcileOutcome.MISSING_IN_SYSTEM,
            "no_system_contribution",
            aggregate(Side.FEED, ("1.25", "EA"), ("2.75", "EA")),
            SideAggregate(side=Side.SYSTEM),
        ),
    )
    rows = (
        unresolved(REASON_MATCH_NEEDS_REVIEW, 2, match_outcome="needs_review"),
        unresolved(REASON_MATCH_UNMATCHED, 3, match_outcome="unmatched"),
        unresolved(REASON_MATCH_INELIGIBLE, 4, match_outcome="ineligible"),
        unresolved(
            REASON_INVALID_AMOUNT, 5, match_outcome="exact", account_ref="ACC-100000"
        ),
        unresolved(
            REASON_INVALID_UNIT, 6, match_outcome="exact", account_ref="ACC-100000"
        ),
        unresolved(
            REASON_INVALID_PERIOD, 7, match_outcome="exact", account_ref="ACC-100001"
        ),
    )
    feed_contributions = sum(len(g.feed.contributions) for g in groups)
    return BatchReconcileResult(
        groups=groups,
        unresolved=rows,
        summary=summarize(
            groups,
            rows,
            feed_rows_in=feed_contributions + len(rows),
            system_rows_in=sum(len(g.system.contributions) for g in groups),
        ),
    )


def colour(value) -> str | None:
    """openpyxl colours as plain text, so a snapshot is JSON-serialisable."""
    return None if value is None else str(value)


def workbook_semantics(content: bytes) -> dict:
    """A canonical snapshot of everything the contract says a workbook holds.

    Deliberately not the raw bytes: the ZIP container carries ordering and
    metadata this project does not control, so determinism is asserted over
    what a reader actually sees.
    """
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    properties = workbook.properties
    snapshot = {
        "properties": {
            "creator": properties.creator,
            "lastModifiedBy": properties.lastModifiedBy,
            "title": properties.title,
            "subject": properties.subject,
            "description": properties.description,
            "created": properties.created.isoformat(),
            "modified": properties.modified.isoformat(),
        },
        "sheet_names": list(workbook.sheetnames),
        "formula_cells": 0,
        "sheets": {},
    }
    formulas = 0
    for name in workbook.sheetnames:
        sheet = workbook[name]
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                fill = colour(cell.fill.fgColor.rgb)
                if cell.value is None and fill in (None, "00000000"):
                    continue
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": cell.value,
                        "data_type": cell.data_type,
                        "bold": bool(cell.font.bold),
                        "font_color": colour(
                            cell.font.color.rgb if cell.font.color else None
                        ),
                        "fill": fill,
                    }
                )
        snapshot["sheets"][name] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "cells": cells,
        }
    snapshot["formula_cells"] = formulas
    workbook.close()
    return snapshot

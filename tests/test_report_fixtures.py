"""T13 against the committed synthetic fixtures, through T9 to T12.

Every expected number is derived from the actual reconciliation result or from
the generator's imported constants — never typed in as an independent business
constant — and each assertion names where its number comes from.
"""

from __future__ import annotations

import hashlib
import importlib.util
import io
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from las.config import RunConfig, load_config
from las.ingest import ACCOUNT_LINK_SCHEMA, ACCOUNT_SCHEMA, feed_schema, read_source
from las.match import build_master_index, match_validated
from las.match import policy_from_config as match_policy
from las.reconcile import (
    BatchReconcileResult,
    FeedBatch,
    ReconcileOutcome,
    reconcile,
)
from las.reconcile import policy_from_config as reconcile_policy
from las.report import (
    OUTCOME_DISPLAY_ORDER,
    OutputFormat,
    ReportModel,
    ReportPolicy,
    build_report,
    policy_from_config,
    render_artifacts,
)
from las.validate import policy_from_config as validation_policy
from las.validate import validate_ingest_result
from reconcile_helpers import ledger_rows_for_fixtures

REPO_ROOT = Path(__file__).resolve().parents[1]
SAMPLES_DIR = REPO_ROOT / "data" / "samples"
FIXTURE_NAMES = ("accounts.csv", "account_links.csv", "feed_alpha.xlsx", "feed_beta.csv")

_spec = importlib.util.spec_from_file_location(
    "generate_fixtures", REPO_ROOT / "scripts" / "generate_fixtures.py"
)
assert _spec is not None and _spec.loader is not None
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)


@pytest.fixture(scope="module")
def config() -> RunConfig:
    return load_config(SAMPLES_DIR / "config.yaml")


@pytest.fixture(scope="module")
def result(config: RunConfig) -> BatchReconcileResult:
    """The committed fixtures carried through every earlier stage."""
    links = read_source(SAMPLES_DIR / "account_links.csv", ACCOUNT_LINK_SCHEMA)
    index = build_master_index(
        read_source(SAMPLES_DIR / "accounts.csv", ACCOUNT_SCHEMA), links
    )
    vpolicy = validation_policy(
        config.validation,
        known_external_refs=frozenset(v for v in links.column("external_ref") if v),
    )
    batches = []
    for feed in config.feeds:
        validated = validate_ingest_result(
            read_source(REPO_ROOT / feed.path, feed_schema(feed)), vpolicy
        )
        batches.append(
            FeedBatch(
                validation=validated,
                matching=match_validated(validated, index, match_policy(config.match)),
            )
        )
    return reconcile(
        batches, ledger_rows_for_fixtures(SAMPLES_DIR, index), reconcile_policy(config.report)
    )


@pytest.fixture(scope="module")
def policy(config: RunConfig) -> ReportPolicy:
    return policy_from_config(config.report)


@pytest.fixture(scope="module")
def report(result: BatchReconcileResult, policy: ReportPolicy) -> ReportModel:
    return build_report(result, policy)


def counts(result: BatchReconcileResult) -> Counter[str]:
    return Counter(group.outcome.value for group in result.groups)


def column(table, index: int, name: str) -> str | None:
    return table.rows[index].texts[table.headers.index(name)]


# --- completeness --------------------------------------------------------------


def test_every_group_appears_in_reconciliation_exactly_once(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    assert report.reconciliation.row_count == len(result.groups)
    assert report.reconciliation.row_count == result.summary.group_count
    keys = [
        (column(report.reconciliation, i, "account_ref"),
         column(report.reconciliation, i, "period"))
        for i in range(report.reconciliation.row_count)
    ]
    assert len(keys) == len(set(keys))


def test_every_unresolved_row_appears_in_the_review_queue_exactly_once(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    assert report.review_queue.row_count == len(result.unresolved)
    identities = [
        (column(report.review_queue, i, "source"),
         column(report.review_queue, i, "source_row"))
        for i in range(report.review_queue.row_count)
    ]
    assert len(identities) == len(set(identities))


def test_the_review_queue_holds_every_unresolved_category(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """Source: the reasons T12 actually recorded, not a typed-in list."""
    rendered = {
        column(report.review_queue, i, "unresolved_reason")
        for i in range(report.review_queue.row_count)
    }
    assert rendered == set(result.summary.unresolved_by_reason)


def test_the_feed_row_accounting_still_holds_through_the_report(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """Source: gen.FEED_ALPHA_ROW_COUNT + gen.FEED_BETA_ROW_COUNT."""
    expected = gen.FEED_ALPHA_ROW_COUNT + gen.FEED_BETA_ROW_COUNT
    accounting = report.summary.accounting
    assert accounting["feed_rows_input"] == expected == result.summary.feed_rows_in
    assert (
        accounting["feed_contributions"] + accounting["unresolved_rows"] == expected
    )


def test_the_system_row_accounting_still_holds_through_the_report(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """Source: gen.LINK_COUNT times the ledger's period count."""
    from las.simulator.store import LEDGER_PERIODS

    expected = gen.LINK_COUNT * len(LEDGER_PERIODS)
    assert report.summary.accounting["system_rows_input"] == expected
    assert report.summary.accounting["system_contributions"] == expected
    assert result.summary.system_rows_in == expected


# --- exceptions ----------------------------------------------------------------


def test_exceptions_holds_exactly_the_four_unsettled_outcomes(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    produced = counts(result)
    expected = sum(
        produced[outcome.value]
        for outcome in (
            ReconcileOutcome.MISMATCH,
            ReconcileOutcome.UNIT_MISMATCH,
            ReconcileOutcome.MISSING_IN_FEED,
            ReconcileOutcome.MISSING_IN_SYSTEM,
        )
    )
    assert report.exceptions.row_count == expected
    assert report.summary.accounting["exception_group_count"] == expected


def test_settled_groups_are_absent_from_exceptions(report: ReportModel) -> None:
    outcomes = {
        column(report.exceptions, i, "outcome")
        for i in range(report.exceptions.row_count)
    }
    assert "balanced" not in outcomes
    assert "within_tolerance" not in outcomes


def test_exceptions_is_the_reconciliation_sheet_filtered(report: ReportModel) -> None:
    full = [
        (column(report.reconciliation, i, "account_ref"),
         column(report.reconciliation, i, "period"))
        for i in range(report.reconciliation.row_count)
    ]
    subset = [
        (column(report.exceptions, i, "account_ref"),
         column(report.exceptions, i, "period"))
        for i in range(report.exceptions.row_count)
    ]
    assert subset == [key for key in full if key in set(subset)]


# --- blanks and exact values ----------------------------------------------------


def test_every_unit_mismatch_group_has_blank_variance_cells(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    rows = [
        i for i in range(report.reconciliation.row_count)
        if column(report.reconciliation, i, "outcome") == "unit_mismatch"
    ]
    assert len(rows) == counts(result)["unit_mismatch"]
    for index in rows:
        for name in ("signed_difference", "absolute_difference", "variance_band_upper"):
            assert column(report.reconciliation, index, name) is None


def test_every_missing_side_group_has_blank_variance_cells(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    produced = counts(result)
    rows = [
        i for i in range(report.reconciliation.row_count)
        if column(report.reconciliation, i, "outcome")
        in ("missing_in_feed", "missing_in_system")
    ]
    assert len(rows) == produced["missing_in_feed"] + produced["missing_in_system"]
    for index in rows:
        for name in ("signed_difference", "absolute_difference", "variance_band_upper"):
            assert column(report.reconciliation, index, name) is None


def test_every_mismatch_keeps_its_exact_decimal_difference(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    by_key = {group.key.sort_key: group for group in result.groups}
    rows = [
        i for i in range(report.reconciliation.row_count)
        if column(report.reconciliation, i, "outcome") == "mismatch"
    ]
    assert len(rows) == counts(result)["mismatch"]
    for index in rows:
        key = (
            column(report.reconciliation, index, "account_ref"),
            column(report.reconciliation, index, "period"),
        )
        group = by_key[key]
        assert Decimal(
            column(report.reconciliation, index, "signed_difference")
        ) == group.signed_difference
        assert column(report.reconciliation, index, "variance_band_upper") is None


def test_every_tolerated_group_keeps_its_selected_band(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    by_key = {group.key.sort_key: group for group in result.groups}
    rows = [
        i for i in range(report.reconciliation.row_count)
        if column(report.reconciliation, i, "outcome") == "within_tolerance"
    ]
    assert len(rows) == counts(result)["within_tolerance"]
    for index in rows:
        key = (
            column(report.reconciliation, index, "account_ref"),
            column(report.reconciliation, index, "period"),
        )
        assert Decimal(
            column(report.reconciliation, index, "variance_band_upper")
        ) == by_key[key].variance_band


def test_the_injected_unit_mismatches_are_visible_in_the_report(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """Source: gen.UNIT_MISMATCH_COUNT — at least that many groups carry two units."""
    multi = [
        i for i in range(report.reconciliation.row_count)
        if (column(report.reconciliation, i, "feed_units") or "").count("|") >= 1
    ]
    assert len(multi) >= gen.UNIT_MISMATCH_COUNT
    for index in multi:
        assert column(report.reconciliation, index, "outcome") == "unit_mismatch"
        assert column(report.reconciliation, index, "feed_total") is None


# --- the summary ----------------------------------------------------------------


def test_the_report_summary_agrees_with_reconciliation(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    produced = counts(result)
    for outcome, count in report.summary.group_outcomes:
        assert count == produced.get(outcome, 0)
    for reason, count in report.summary.group_reasons:
        assert count == result.summary.groups_by_reason.get(reason, 0)
    for reason, count in report.summary.unresolved_reasons:
        assert count == result.summary.unresolved_by_reason.get(reason, 0)


def test_zero_count_rows_are_still_stated(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """The fixtures produce no balanced group; the report says so out loud."""
    assert counts(result)["balanced"] == 0
    stated = dict(report.summary.group_outcomes)
    assert stated["balanced"] == 0
    assert [key for key, _ in report.summary.group_outcomes] == [
        outcome.value for outcome in OUTCOME_DISPLAY_ORDER
    ]


def test_the_band_counts_add_up_to_the_banded_groups(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    banded = sum(1 for group in result.groups if group.variance_band is not None)
    assert sum(count for _, count in report.summary.variance_bands) == banded


# --- artifacts over the real data ------------------------------------------------


def test_all_three_artifacts_render_from_the_committed_fixtures(
    result: BatchReconcileResult, policy: ReportPolicy
) -> None:
    artifacts = render_artifacts(result, policy)
    assert [a.output_format for a in artifacts] == [
        OutputFormat.XLSX, OutputFormat.JSON, OutputFormat.HTML
    ]
    assert all(a.size > 0 for a in artifacts)


def test_the_workbook_over_real_data_carries_no_formula(
    result: BatchReconcileResult, policy: ReportPolicy
) -> None:
    content = render_artifacts(result, policy)[0].content
    book = load_workbook(io.BytesIO(content), data_only=False)
    assert book.sheetnames == ["Summary", "Exceptions", "Reconciliation", "Review Queue"]
    assert sum(
        1 for sheet in book for row in sheet.iter_rows() for cell in row
        if cell.data_type == "f"
    ) == 0


def test_the_workbook_dimensions_match_the_model(
    report: ReportModel, result: BatchReconcileResult, policy: ReportPolicy
) -> None:
    content = render_artifacts(result, policy)[0].content
    book = load_workbook(io.BytesIO(content))
    for table in report.tables:
        assert book[table.sheet.value].max_row == table.row_count + 1


def test_reporting_changed_no_reconciliation_outcome(
    result: BatchReconcileResult, policy: ReportPolicy
) -> None:
    before = [
        (g.outcome, g.reason, g.signed_difference, g.variance_band) for g in result.groups
    ]
    render_artifacts(result, policy)
    after = [
        (g.outcome, g.reason, g.signed_difference, g.variance_band) for g in result.groups
    ]
    assert before == after


def test_reporting_changed_no_unresolved_row(
    result: BatchReconcileResult, policy: ReportPolicy
) -> None:
    before = [(r.reason, r.source, r.source_row, r.match_outcome) for r in result.unresolved]
    render_artifacts(result, policy)
    assert [
        (r.reason, r.source, r.source_row, r.match_outcome) for r in result.unresolved
    ] == before


def test_duplicate_contributions_remain_represented_through_the_counts(
    report: ReportModel, result: BatchReconcileResult
) -> None:
    """T12 keeps every duplicate; the report states them, never collapses them."""
    total = sum(
        int(column(report.reconciliation, i, "feed_contribution_count"))
        for i in range(report.reconciliation.row_count)
    )
    assert total == result.summary.feed_contributions
    assert total == report.summary.accounting["feed_contributions"]


# --- the fixtures are untouched ---------------------------------------------------


def test_reporting_leaves_every_committed_fixture_byte_identical() -> None:
    documented = {
        match.group(2).split("/")[-1]: match.group(1)
        for match in re.finditer(
            r"^([0-9a-f]{64})\s+(data/samples/\S+)$",
            (REPO_ROOT / "docs" / "DEVELOPMENT_PLAN.md").read_text(encoding="utf-8"),
            re.MULTILINE,
        )
    }
    assert set(documented) == set(FIXTURE_NAMES)
    actual = {
        name: hashlib.sha256((SAMPLES_DIR / name).read_bytes()).hexdigest()
        for name in FIXTURE_NAMES
    }
    assert actual == documented

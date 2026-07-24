"""Fixture-generator behavior: determinism, schemas, defects, and safety."""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import re
from difflib import SequenceMatcher
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_PATH = REPO_ROOT / "scripts" / "generate_fixtures.py"

_spec = importlib.util.spec_from_file_location("generate_fixtures", SCRIPT_PATH)
assert _spec is not None and _spec.loader is not None
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)

EXPECTED_FILES = {
    "accounts.csv",
    "account_links.csv",
    "feed_alpha.xlsx",
    "feed_beta.csv",
}

ACCOUNT_REF_PATTERN = re.compile(r"^ACC-\d{6}$")
EXTERNAL_REF_PATTERN = re.compile(r"^EXT-\d{5}$")
RECORD_ID_PATTERN = re.compile(r"^OPR-\d{8}$")

FORBIDDEN_FRAGMENTS = ("http", "://", "www.", ":\\", "\\\\")
FORBIDDEN_TERMS = re.compile(
    r"\b(invoice|bonus|beverage|beer|brewery|wholesaler|sap)\b", re.IGNORECASE
)


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, dict]:
    """One shared default-seed generation used by the read-only tests."""
    target = tmp_path_factory.mktemp("fixtures")
    summary = gen.generate_fixtures(output_dir=target, seed=gen.DEFAULT_SEED)
    return target, summary


def read_csv(path: Path, delimiter: str = ",", encoding: str = "utf-8") -> list[dict[str, str]]:
    with path.open(encoding=encoding, newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def read_alpha(directory: Path) -> list[dict[str, str]]:
    sheet = load_workbook(directory / "feed_alpha.xlsx").active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(cell) for cell in rows[0]]
    return [
        {
            header: "" if cell is None else str(cell)
            for header, cell in zip(headers, row, strict=True)
        }
        for row in rows[1:]
    ]


def read_beta(directory: Path) -> list[dict[str, str]]:
    return read_csv(directory / "feed_beta.csv", delimiter=";", encoding="utf-8-sig")


def file_hashes(directory: Path) -> dict[str, str]:
    return {
        name: hashlib.sha256((directory / name).read_bytes()).hexdigest()
        for name in sorted(EXPECTED_FILES)
    }


# --- determinism -------------------------------------------------------------


def test_same_seed_produces_byte_identical_files(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    gen.generate_fixtures(output_dir=first, seed=gen.DEFAULT_SEED)
    gen.generate_fixtures(output_dir=second, seed=gen.DEFAULT_SEED)
    assert file_hashes(first) == file_hashes(second)


def test_different_seeds_produce_different_content(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    gen.generate_fixtures(output_dir=first, seed=gen.DEFAULT_SEED)
    gen.generate_fixtures(output_dir=second, seed=gen.DEFAULT_SEED + 1)
    first_hashes = file_hashes(first)
    second_hashes = file_hashes(second)
    assert first_hashes["accounts.csv"] != second_hashes["accounts.csv"]
    assert first_hashes["feed_alpha.xlsx"] != second_hashes["feed_alpha.xlsx"]
    assert first_hashes["feed_beta.csv"] != second_hashes["feed_beta.csv"]


def test_regeneration_leaves_unrelated_files_untouched(tmp_path: Path) -> None:
    bystander = tmp_path / "config.yaml"
    bystander.write_bytes(b"seed: 42\n# unrelated sample configuration\n")
    gen.generate_fixtures(output_dir=tmp_path, seed=gen.DEFAULT_SEED)
    assert bystander.read_bytes() == b"seed: 42\n# unrelated sample configuration\n"


# --- files, schemas, and row counts ------------------------------------------


def test_expected_files_are_created(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    assert {path.name for path in directory.iterdir()} == EXPECTED_FILES


def test_row_counts_match_summary_and_constants(generated: tuple[Path, dict]) -> None:
    directory, summary = generated
    accounts = read_csv(directory / "accounts.csv")
    links = read_csv(directory / "account_links.csv")
    alpha = read_alpha(directory)
    beta = read_beta(directory)

    assert len(accounts) == gen.ACCOUNT_COUNT == summary["files"]["accounts.csv"]["rows"]
    assert len(links) == gen.LINK_COUNT == summary["files"]["account_links.csv"]["rows"]
    assert len(alpha) == gen.FEED_ALPHA_ROW_COUNT == summary["files"]["feed_alpha.xlsx"]["rows"]
    assert len(beta) == gen.FEED_BETA_ROW_COUNT == summary["files"]["feed_beta.csv"]["rows"]


def test_schemas_use_the_documented_headers(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    with (directory / "accounts.csv").open(encoding="utf-8", newline="") as handle:
        assert next(csv.reader(handle)) == gen.ACCOUNT_HEADERS
    with (directory / "account_links.csv").open(encoding="utf-8", newline="") as handle:
        assert next(csv.reader(handle)) == gen.LINK_HEADERS
    sheet = load_workbook(directory / "feed_alpha.xlsx").active
    assert [cell.value for cell in sheet[1]] == gen.FEED_ALPHA_HEADERS
    with (directory / "feed_beta.csv").open(encoding="utf-8-sig", newline="") as handle:
        assert next(csv.reader(handle, delimiter=";")) == gen.FEED_BETA_HEADERS


# --- synthetic-only content ---------------------------------------------------


def test_identifier_patterns(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    accounts = read_csv(directory / "accounts.csv")
    links = read_csv(directory / "account_links.csv")

    assert all(ACCOUNT_REF_PATTERN.fullmatch(a["account_ref"]) for a in accounts)
    assert all(EXTERNAL_REF_PATTERN.fullmatch(link["external_ref"]) for link in links)
    assert all(ACCOUNT_REF_PATTERN.fullmatch(link["account_ref"]) for link in links)

    alpha_ids = [row["Record Id"] for row in read_alpha(directory)]
    assert all(RECORD_ID_PATTERN.fullmatch(record_id) for record_id in alpha_ids)

    beta_ids = [row["REC_ID"] for row in read_beta(directory)]
    invalid = [record_id for record_id in beta_ids if not RECORD_ID_PATTERN.fullmatch(record_id)]
    assert sorted(invalid) == sorted(gen.INVALID_RECORD_ID_VALUES)


def test_all_emails_use_example_domain(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    for account in read_csv(directory / "accounts.csv"):
        assert account["contact_email"].endswith("@example.com")


ACCOUNT_NAME_PATTERN = re.compile(r"^(Demo Account \d{6}|Synthetic Twin Account \d{6})$")
STREET_PATTERN = re.compile(r"^\d{6} Example Avenue$")
CITY_PATTERN = re.compile(r"^Synthetic City \d{2}$")
POSTAL_PATTERN = re.compile(r"^\d{5}$")
REALISTIC_SUFFIXES = re.compile(
    r"\b(LLC|Ltd|Inc|PLC|GmbH|Corp|Company|Group|Sons|Holdings)\b", re.IGNORECASE
)


def test_account_names_follow_synthetic_convention(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    for account in read_csv(directory / "accounts.csv"):
        assert ACCOUNT_NAME_PATTERN.fullmatch(account["account_name"]), account["account_name"]


def test_addresses_follow_synthetic_convention(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    for account in read_csv(directory / "accounts.csv"):
        assert STREET_PATTERN.fullmatch(account["street"]), account["street"]
        assert CITY_PATTERN.fullmatch(account["city"]), account["city"]
        assert POSTAL_PATTERN.fullmatch(account["postal_code"]), account["postal_code"]


def test_no_personal_names_or_realistic_company_suffixes(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    names = [account["account_name"] for account in read_csv(directory / "accounts.csv")]
    names += [row["Reported By"] for row in read_alpha(directory)]
    names += [row["SRC_NAME"] for row in read_beta(directory)]
    for name in names:
        assert not REALISTIC_SUFFIXES.search(name), name
        assert (
            name == ""
            or name.startswith(("Demo Account ", "Synthetic Twin Account ", "Unlisted Source "))
        ), name


def test_no_forbidden_content_in_any_generated_value(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    all_text: list[str] = []
    for name in ("accounts.csv", "account_links.csv"):
        all_text.append((directory / name).read_text(encoding="utf-8"))
    all_text.append((directory / "feed_beta.csv").read_text(encoding="utf-8-sig"))
    for row in read_alpha(directory):
        all_text.extend(row.values())

    combined = "\n".join(all_text)
    for fragment in FORBIDDEN_FRAGMENTS:
        assert fragment not in combined, f"forbidden fragment {fragment!r} found"
    assert not FORBIDDEN_TERMS.search(combined)
    for domain in re.findall(r"@([\w.-]+)", combined):
        assert domain == "example.com"


# --- defect injection ---------------------------------------------------------


def test_summary_defect_counts_are_exact(generated: tuple[Path, dict]) -> None:
    _, summary = generated
    assert summary["defects"] == {
        "duplicate_records": 4,
        "invalid_record_id_format": 3,
        "orphan_external_refs": 5,
        "missing_required_values": 4,
        "invalid_amounts": 3,
        "unit_mismatches": 3,
        "ambiguous_matches": 2,
    }
    assert summary["seed"] == gen.DEFAULT_SEED


def test_injected_duplicates_are_genuinely_duplicated(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    rows = [tuple(row.items()) for row in read_alpha(directory)]
    duplicated = {row for row in rows if rows.count(row) == 2}
    assert len(duplicated) == gen.DUPLICATE_RECORD_COUNT


def test_orphan_references_are_absent_from_mapping(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    mapped = {link["external_ref"] for link in read_csv(directory / "account_links.csv")}
    feed_refs = {
        row["External Reference"] for row in read_alpha(directory) if row["External Reference"]
    } | {row["EXT_REF"] for row in read_beta(directory) if row["EXT_REF"]}
    orphans = {ref for ref in feed_refs if EXTERNAL_REF_PATTERN.fullmatch(ref)} - mapped
    assert len(orphans) == gen.ORPHAN_REFERENCE_COUNT
    assert all(ref.startswith("EXT-9") for ref in orphans)


def test_missing_required_values_count(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    alpha_missing = [r for r in read_alpha(directory) if r["Quantity"] == ""]
    beta_missing = [r for r in read_beta(directory) if r["AMT"] == ""]
    assert len(alpha_missing) + len(beta_missing) == gen.MISSING_VALUE_COUNT
    assert len(alpha_missing) == len(beta_missing)


def test_invalid_amounts_count(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    invalid = 0
    for row in read_alpha(directory):
        text = row["Quantity"]
        if text == "":
            continue
        try:
            if float(text) < 0:
                invalid += 1
        except ValueError:
            invalid += 1
    assert invalid == gen.INVALID_AMOUNT_COUNT


def test_unit_mismatches_between_feeds(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    alpha_units = {row["Record Id"]: row["Quantity Unit"] for row in read_alpha(directory)}
    beta_units = {row["REC_ID"]: row["UOM"] for row in read_beta(directory)}
    shared = set(alpha_units) & set(beta_units)
    mismatched = {rid for rid in shared if alpha_units[rid] != beta_units[rid]}
    assert len(mismatched) == gen.UNIT_MISMATCH_COUNT


def test_ambiguous_fixtures_are_actually_ambiguous(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    ambiguous_rows = [
        row
        for row in read_alpha(directory)
        if row["External Reference"] == "" and row["Reported By"]
    ]
    assert len(ambiguous_rows) == gen.AMBIGUOUS_MATCH_COUNT

    twins = gen.AMBIGUOUS_ACCOUNT_NAMES
    account_names = {a["account_name"] for a in read_csv(directory / "accounts.csv")}
    assert set(twins) <= account_names

    for row in ambiguous_rows:
        source = row["Reported By"]
        scores = [SequenceMatcher(None, source, twin).ratio() for twin in twins]
        assert all(score >= 0.85 for score in scores)
        assert abs(scores[0] - scores[1]) <= 0.05


# --- xlsx determinism details -------------------------------------------------


def test_xlsx_metadata_is_neutral_and_fixed(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    properties = load_workbook(directory / "feed_alpha.xlsx").properties
    assert properties.creator == gen.XLSX_CREATOR
    assert properties.lastModifiedBy == gen.XLSX_CREATOR
    assert properties.created == gen.XLSX_TIMESTAMP
    assert properties.modified == gen.XLSX_TIMESTAMP


def test_xlsx_zip_entries_have_fixed_timestamps(generated: tuple[Path, dict]) -> None:
    directory, _ = generated
    with ZipFile(directory / "feed_alpha.xlsx") as archive:
        for info in archive.infolist():
            assert info.date_time == (1980, 1, 1, 0, 0, 0), info.filename
            assert info.create_system == 3


def test_xlsx_document_xml_carries_no_identifying_metadata(
    generated: tuple[Path, dict]
) -> None:
    import getpass
    import platform

    directory, _ = generated
    username = getpass.getuser()
    machine = platform.node()

    with ZipFile(directory / "feed_alpha.xlsx") as archive:
        names = archive.namelist()
        core_xml = archive.read("docProps/core.xml").decode("utf-8")
        app_xml = archive.read("docProps/app.xml").decode("utf-8") if (
            "docProps/app.xml" in names
        ) else ""

    for document in (core_xml, app_xml):
        if username:
            assert username not in document
        if machine:
            assert machine not in document
        assert ":\\" not in document and "\\\\" not in document  # no absolute/UNC paths
        assert "Company" not in document or "<Company/>" in document or (
            "<Company></Company>" in document
        )

    # Only the generic generator identity and the fixed timestamp appear.
    assert core_xml.count(gen.XLSX_CREATOR) == 2  # creator + lastModifiedBy
    fixed_stamp = gen.XLSX_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    assert core_xml.count(fixed_stamp) == 2  # created + modified
    for match in re.findall(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", core_xml):
        assert match == fixed_stamp


# --- CLI ----------------------------------------------------------------------


def test_cli_prints_machine_readable_summary(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    import json

    exit_code = gen.main(["--seed", "7", "--output-dir", str(tmp_path / "out")])
    assert exit_code == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["seed"] == 7
    assert set(summary["files"]) == EXPECTED_FILES

"""T9 XLSX reader: worksheets, cell types, formulas, provenance, purity."""

from __future__ import annotations

import datetime as dt
import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook

from ingest_helpers import schema
from las.ingest import (
    ExtraColumnPolicy,
    FormulaCellError,
    HeaderResolutionError,
    IngestOutcome,
    MalformedRowError,
    SourceFormat,
    SourceReadError,
    UnsupportedCellValueError,
    WorksheetSelectionError,
    cell_to_text,
    read_source,
    read_xlsx,
)

HEADERS = ["record_id", "external_ref", "amount", "unit"]
ROWS = [
    ["OPR-00000001", "EXT-10000", "10.50", "EA"],
    ["OPR-00000002", "EXT-10001", "20.25", "KG"],
]


def build(
    tmp_path: Path,
    rows: list[list],
    *,
    headers: list | None = None,
    name: str = "probe.xlsx",
    sheet_title: str = "records",
    extra_sheets: tuple[str, ...] = (),
    cells: dict[str, object] | None = None,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    if headers is not None:
        sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for reference, value in (cells or {}).items():
        sheet[reference] = value
    for title in extra_sheets:
        workbook.create_sheet(title)
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


# --- ordinary reads -----------------------------------------------------------


def test_ordinary_valid_workbook(tmp_path: Path) -> None:
    result = read_xlsx(build(tmp_path, ROWS, headers=HEADERS), schema())

    assert result.source == "probe.xlsx"
    assert result.source_format is SourceFormat.XLSX
    assert result.worksheet == "records"
    assert result.headers == ("record_id", "external_ref", "amount", "unit")
    assert result.header_row == 1
    assert result.row_count == 2
    assert result.skipped_blank_rows == 0
    assert result.outcome is IngestOutcome.ROWS
    assert result.column("record_id") == ("OPR-00000001", "OPR-00000002")
    assert [row.provenance.source_row for row in result.rows] == [2, 3]
    assert result.rows[0].provenance.worksheet == "records"


def test_alias_headers_resolve(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "EXT-10000", "1", "EA"]],
        headers=["Record Id", "External Reference", "Quantity", "Quantity Unit"],
    )
    result = read_xlsx(path, schema())
    assert result.headers == ("record_id", "external_ref", "amount", "unit")


def test_read_source_dispatches_to_the_xlsx_reader(tmp_path: Path) -> None:
    result = read_source(build(tmp_path, ROWS, headers=HEADERS), schema())
    assert result.source_format is SourceFormat.XLSX


def test_source_label_overrides_the_basename(tmp_path: Path) -> None:
    path = build(tmp_path, ROWS, headers=HEADERS)
    result = read_xlsx(path, schema(), source_label="feed_probe")
    assert result.source == "feed_probe"


# --- worksheet selection ------------------------------------------------------


def test_default_worksheet_is_the_first_in_workbook_order(tmp_path: Path) -> None:
    path = build(tmp_path, ROWS, headers=HEADERS, sheet_title="first", extra_sheets=("second",))
    result = read_xlsx(path, schema())
    assert result.worksheet == "first"


def test_explicit_worksheet_selection(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "decoy"
    workbook.active.append(["nonsense"])
    target = workbook.create_sheet("wanted")
    target.append(HEADERS)
    target.append(ROWS[0])
    path = tmp_path / "probe.xlsx"
    workbook.save(path)
    workbook.close()

    result = read_xlsx(path, schema(), worksheet="wanted")
    assert result.worksheet == "wanted"
    assert result.row_count == 1


def test_missing_worksheet_is_a_typed_error_listing_what_exists(tmp_path: Path) -> None:
    path = build(tmp_path, ROWS, headers=HEADERS, sheet_title="records")
    with pytest.raises(WorksheetSelectionError, match="no such worksheet") as excinfo:
        read_xlsx(path, schema(), worksheet="absent")
    assert "records" in str(excinfo.value)  # the caller is told what is available
    assert excinfo.value.worksheet == "absent"


def test_missing_file_is_a_typed_error(tmp_path: Path) -> None:
    with pytest.raises(SourceReadError, match="does not exist") as excinfo:
        read_xlsx(tmp_path / "absent.xlsx", schema())
    assert excinfo.value.source == "absent.xlsx"


def test_non_workbook_content_is_a_typed_error(tmp_path: Path) -> None:
    path = tmp_path / "probe.xlsx"
    path.write_bytes(b"this is definitely not a workbook")
    with pytest.raises(SourceReadError, match="could not be opened"):
        read_xlsx(path, schema())


# --- blank rows, empty cells, header-only ------------------------------------


def test_blank_rows_are_skipped_and_counted_with_honest_row_numbers(
    tmp_path: Path,
) -> None:
    rows = [ROWS[0], [None, None, None, None], [None, None, None, None], ROWS[1]]
    result = read_xlsx(build(tmp_path, rows, headers=HEADERS), schema())
    assert result.row_count == 2
    assert result.skipped_blank_rows == 2
    # Rows 3 and 4 were blank, so the second record really is on row 5.
    assert [row.provenance.source_row for row in result.rows] == [2, 5]


def test_absent_cells_become_missing_values(tmp_path: Path) -> None:
    result = read_xlsx(
        build(tmp_path, [["OPR-00000001", None, None, None]], headers=HEADERS), schema()
    )
    row = result.rows[0]
    assert row.get("record_id") == "OPR-00000001"
    assert row.get("external_ref") is None
    assert row.get("amount") is None
    assert row.get("unit") is None


def test_empty_string_cell_is_also_missing(tmp_path: Path) -> None:
    result = read_xlsx(
        build(tmp_path, [["OPR-00000001", "", "", ""]], headers=HEADERS), schema()
    )
    assert result.rows[0].get("external_ref") is None


def test_header_only_worksheet_is_a_deterministic_outcome(tmp_path: Path) -> None:
    result = read_xlsx(build(tmp_path, [], headers=HEADERS), schema())
    assert result.rows == ()
    assert result.outcome is IngestOutcome.HEADER_ONLY
    assert result.headers == ("record_id", "external_ref", "amount", "unit")


def test_blank_rows_before_the_header_are_skipped_not_counted(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "records"
    sheet.append([None, None])
    sheet.append(HEADERS)
    sheet.append(ROWS[0])
    path = tmp_path / "probe.xlsx"
    workbook.save(path)
    workbook.close()

    result = read_xlsx(path, schema())
    assert result.header_row == 2
    assert result.skipped_blank_rows == 0
    assert result.rows[0].provenance.source_row == 3


def test_worksheet_with_no_header_row_is_rejected(tmp_path: Path) -> None:
    path = build(tmp_path, [], headers=None)
    with pytest.raises(SourceReadError, match="no header row"):
        read_xlsx(path, schema())


# --- header problems ----------------------------------------------------------


def test_duplicate_header_is_rejected(tmp_path: Path) -> None:
    path = build(
        tmp_path, [ROWS[0]], headers=["record_id", "external_ref", "amount", "record_id"]
    )
    with pytest.raises(HeaderResolutionError, match="repeats"):
        read_xlsx(path, schema())


def test_missing_required_header_is_rejected(tmp_path: Path) -> None:
    path = build(tmp_path, [["OPR-00000001"]], headers=["record_id"])
    with pytest.raises(HeaderResolutionError, match="absent"):
        read_xlsx(path, schema())


def test_blank_header_cell_between_columns_is_rejected(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "x", "EXT-10000", "1"]],
        headers=["record_id", None, "external_ref", "amount"],
    )
    with pytest.raises(HeaderResolutionError, match="blank") as excinfo:
        read_xlsx(path, schema())
    assert excinfo.value.column == 2


def test_trailing_blank_header_cells_are_trimmed(tmp_path: Path) -> None:
    """A header row padded out by the format declares only its real columns."""
    path = build(tmp_path, [ROWS[0]], headers=[*HEADERS, None, None])
    result = read_xlsx(path, schema())
    assert result.headers == ("record_id", "external_ref", "amount", "unit")
    assert result.row_count == 1


# --- row width ----------------------------------------------------------------


def test_value_beyond_the_header_span_is_a_structural_error(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [ROWS[0]],
        headers=["record_id", "external_ref", "amount", "unit"],
        cells={"F2": "stray"},
    )
    with pytest.raises(MalformedRowError, match="beyond the 4 column") as excinfo:
        read_xlsx(path, schema())
    assert excinfo.value.row == 2
    assert excinfo.value.column == 6


def test_structural_error_does_not_leak_the_offending_value(tmp_path: Path) -> None:
    path = build(
        tmp_path, [ROWS[0]], headers=HEADERS, cells={"F2": "SENSITIVE-STRAY-VALUE"}
    )
    with pytest.raises(MalformedRowError) as excinfo:
        read_xlsx(path, schema())
    assert "SENSITIVE-STRAY-VALUE" not in str(excinfo.value)


def test_short_logical_rows_are_treated_as_missing_trailing_cells(
    tmp_path: Path,
) -> None:
    """The format pads rows, so a short row is missing values, not malformed."""
    path = build(tmp_path, [["OPR-00000001", "EXT-10000"]], headers=HEADERS)
    result = read_xlsx(path, schema())
    assert result.row_count == 1
    assert result.rows[0].get("amount") is None
    assert result.rows[0].get("unit") is None


# --- cell value policy --------------------------------------------------------


def test_numeric_boolean_and_date_cells_normalize_deterministically(
    tmp_path: Path,
) -> None:
    path = build(
        tmp_path,
        [
            ["OPR-00000001", "EXT-10000", 12.5, "EA"],
            ["OPR-00000002", "EXT-10001", 52, "KG"],
            ["OPR-00000003", "EXT-10002", True, dt.datetime(2026, 1, 2, 3, 4)],
        ],
        headers=HEADERS,
    )
    result = read_xlsx(path, schema())
    assert result.column("amount") == ("12.5", "52", "TRUE")
    assert result.rows[2].get("unit") == "2026-01-02T03:04:00"


def test_numeric_values_are_neither_padded_nor_text_reformatted(
    tmp_path: Path,
) -> None:
    """A number keeps its own shortest form; text keeps its exact characters.

    The pair documents the one genuine cross-format ambiguity: a trailing
    zero exists in text but not in the numeric domain, so ingestion neither
    invents one for the number nor strips one from the text.
    """
    path = build(
        tmp_path,
        [
            ["OPR-00000001", "EXT-10000", 12.50, "EA"],
            ["OPR-00000002", "EXT-10001", "12.50", "EA"],
        ],
        headers=HEADERS,
    )
    result = read_xlsx(path, schema())
    assert result.column("amount") == ("12.5", "12.50")


def test_formula_cells_are_refused_rather_than_trusted(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "EXT-10000", "=1+1", "EA"]],
        headers=HEADERS,
    )
    with pytest.raises(FormulaCellError, match="never evaluates formulas") as excinfo:
        read_xlsx(path, schema())
    assert excinfo.value.row == 2
    assert excinfo.value.column == 3
    assert excinfo.value.worksheet == "records"


def test_formula_in_the_header_row_is_also_refused(tmp_path: Path) -> None:
    path = build(
        tmp_path, [ROWS[0]], headers=["record_id", "=A1", "amount", "unit"]
    )
    with pytest.raises(FormulaCellError) as excinfo:
        read_xlsx(path, schema())
    assert excinfo.value.row == 1


def test_unsupported_cell_type_is_refused_by_the_allowlist(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "EXT-10000", dt.time(3, 4), "EA"]],
        headers=HEADERS,
    )
    with pytest.raises(UnsupportedCellValueError, match="'time' is not supported") as excinfo:
        read_xlsx(path, schema())
    assert excinfo.value.row == 2
    assert excinfo.value.column == 3


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        ("text", "text"),
        ("  spaced  ", "  spaced  "),
        (0, "0"),
        (-7, "-7"),
        (12.5, "12.5"),
        (0.1 + 0.2, "0.30000000000000004"),
        (True, "TRUE"),
        (False, "FALSE"),
        (dt.date(2026, 5, 6), "2026-05-06"),
        (dt.datetime(2026, 5, 6, 7, 8, 9), "2026-05-06T07:08:09"),
    ],
)
def test_cell_value_allowlist_normalization(value: object, expected: str | None) -> None:
    assert cell_to_text(value, operation="probe", source="probe.xlsx") == expected


@pytest.mark.parametrize("value", [dt.time(1, 2), object(), [1], {"a": 1}, (1,), {1}])
def test_cell_values_outside_the_allowlist_are_refused(value: object) -> None:
    with pytest.raises(UnsupportedCellValueError, match="not supported"):
        cell_to_text(value, operation="probe", source="probe.xlsx")


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_non_finite_numbers_are_refused(value: float) -> None:
    with pytest.raises(UnsupportedCellValueError, match="non-finite"):
        cell_to_text(value, operation="probe", source="probe.xlsx")


def test_unsupported_cell_error_names_the_type_not_the_value() -> None:
    class Secretive:
        def __repr__(self) -> str:  # pragma: no cover - must never run
            raise AssertionError("the value must never be serialized")

        def __str__(self) -> str:  # pragma: no cover - must never run
            raise AssertionError("the value must never be serialized")

    with pytest.raises(UnsupportedCellValueError, match="Secretive"):
        cell_to_text(Secretive(), operation="probe", source="probe.xlsx")


# --- extra columns ------------------------------------------------------------


def test_extra_column_is_retained_and_reported(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "EXT-10000", "1", "EA", "Demo Account 000001"]],
        headers=[*HEADERS, "Reported By"],
    )
    result = read_xlsx(path, schema())
    assert result.unmapped_headers == ("reported by",)
    assert result.rows[0].extras == {"reported by": "Demo Account 000001"}


def test_extra_column_can_be_reported_but_dropped(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [["OPR-00000001", "EXT-10000", "1", "EA", "Demo Account 000001"]],
        headers=[*HEADERS, "Reported By"],
    )
    result = read_xlsx(path, schema(extra_columns=ExtraColumnPolicy.IGNORE))
    assert result.unmapped_headers == ("reported by",)
    assert result.rows[0].extras == {}


# --- business defects survive -------------------------------------------------


def test_business_invalid_values_are_preserved_for_validation(tmp_path: Path) -> None:
    path = build(
        tmp_path,
        [
            ["OPR12345678", "EXT-10000", "not-a-number", "EA"],
            ["opr-00000001", "EXT-90000", -42.5, "PALLET"],
            ["OPR-00000003", None, "12.34.56", None],
        ],
        headers=HEADERS,
    )
    result = read_xlsx(path, schema())
    assert result.column("record_id") == ("OPR12345678", "opr-00000001", "OPR-00000003")
    assert result.column("amount") == ("not-a-number", "-42.5", "12.34.56")
    assert result.column("unit") == ("EA", "PALLET", None)
    assert result.column("external_ref") == ("EXT-10000", "EXT-90000", None)


# --- purity and determinism ---------------------------------------------------


def test_repeated_reads_are_identical(tmp_path: Path) -> None:
    path = build(tmp_path, ROWS, headers=HEADERS)
    assert read_xlsx(path, schema()) == read_xlsx(path, schema())


def test_reading_leaves_the_workbook_byte_identical(tmp_path: Path) -> None:
    path = build(tmp_path, ROWS, headers=HEADERS)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    read_xlsx(path, schema())
    read_xlsx(path, schema())
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_errors_never_expose_a_filesystem_path(tmp_path: Path) -> None:
    path = build(
        tmp_path, [["OPR-00000001", "EXT-10000", "=1+1", "EA"]], headers=HEADERS
    )
    with pytest.raises(FormulaCellError) as excinfo:
        read_xlsx(path, schema())
    message = str(excinfo.value)
    assert str(tmp_path) not in message
    assert "probe.xlsx" in message


def test_workbook_without_worksheets_is_a_typed_error() -> None:
    """openpyxl cannot save such a file, so the guard is unit-tested directly.

    Without it, worksheet selection would raise a bare IndexError instead of
    the package's own typed error.
    """
    from las.ingest.readers import _select_worksheet

    class Sheetless:
        worksheets: list[object] = []
        sheetnames: list[str] = []

    with pytest.raises(WorksheetSelectionError, match="no worksheets"):
        _select_worksheet(Sheetless(), None, operation="read_xlsx", source="probe.xlsx")

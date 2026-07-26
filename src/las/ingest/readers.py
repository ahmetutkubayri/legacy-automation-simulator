"""Read-only CSV and XLSX readers producing identical canonical results.

Both readers open their source for reading only. Nothing here writes, and
the XLSX reader never saves a workbook or touches its properties, so a
source file is byte-identical before and after ingestion.

What counts as a structural failure
-----------------------------------
Ingestion rejects problems with the *file or its schema*: an unreadable or
undecodable source, no header row, a blank or duplicated header, two columns
claiming one canonical field, a missing required field, a row whose shape
contradicts the header row, a formula cell, and a cell type outside the safe
allowlist.

Ingestion never judges *content*. Malformed identifiers, empty values,
unparseable numbers, unknown units, orphan references and ambiguous names
all pass through exactly as written, because deciding they are wrong is the
validation stage's job and it needs to see them to do it.

CSV specifics
-------------
Encoding is resolved through a fixed fallback chain,
:data:`DEFAULT_ENCODINGS`. ``utf-8-sig`` comes first and transparently
handles UTF-8 both with and without a byte-order mark; ``cp1252`` is the
last resort for legacy single-byte files. The first codec that decodes the
whole file wins, so the choice never depends on locale.

The delimiter is inferred deterministically, not sniffed heuristically: each
candidate in :data:`DELIMITER_CANDIDATES` parses the header line, and the
one yielding the most fields wins, with ties going to the earliest
candidate. An explicit ``delimiter`` argument skips inference entirely.

XLSX specifics
--------------
Opened read-only with ``data_only=False``, so cached formula results are
never read as if they were data, and with ``keep_links=False`` and
``keep_vba=False`` so external workbook links and macros are not retained.
Nothing is evaluated: not formulas, not macros, not external content.
Worksheet selection is explicit, or else the first worksheet in workbook
order.

Every cell inside the observable row span is put through the formula and
allowlist checks *before* any column is discarded, so an ignored or
unmapped column cannot become a hiding place for a formula or an
unsupported value.

Because the format pads every row out to the sheet's widest column, a
"short" row is indistinguishable from one with trailing empty cells; such
cells are simply missing values. A value *beyond* the header row's span is a
different matter — that is data outside the declared columns, and it is a
structural error rather than something to silently truncate.
"""

from __future__ import annotations

import csv
import errno
from collections.abc import Iterable, Sequence
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from las.ingest.errors import (
    FormulaCellError,
    IngestError,
    MalformedRowError,
    SourceReadError,
    UnsupportedFormatError,
    WorksheetSelectionError,
)
from las.ingest.headers import (
    ExtraColumnPolicy,
    HeaderResolution,
    IngestSchema,
    resolve_headers,
)
from las.ingest.models import (
    IngestedRow,
    IngestOutcome,
    IngestResult,
    RowProvenance,
    SourceFormat,
    safe_source_label,
)
from las.ingest.values import cell_to_text

#: Ordered encoding fallback chain. ``utf-8-sig`` also decodes BOM-less UTF-8.
DEFAULT_ENCODINGS = ("utf-8-sig", "cp1252")

#: Ordered delimiter candidates; earliest wins a tie on field count.
DELIMITER_CANDIDATES = (",", ";", "\t", "|")

#: openpyxl's data-type marker for a formula cell.
_FORMULA_DATA_TYPE = "f"


def _decode(
    path: Path, encodings: Sequence[str], *, operation: str, source: str
) -> tuple[str, str]:
    """Return the decoded text and the codec that succeeded."""
    if not encodings:
        raise SourceReadError(
            "no candidate encodings were supplied", operation=operation, source=source
        )
    try:
        raw = path.read_bytes()
    except OSError as exc:
        # The errno symbol is stable and locale-independent; the operating
        # system's own message is neither, and can quote the full path.
        symbol = errno.errorcode.get(exc.errno, "unknown")
        raise SourceReadError(
            f"the source could not be read ({symbol})",
            operation=operation,
            source=source,
        ) from None
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise SourceReadError(
        f"the source did not decode as any of: {', '.join(encodings)}",
        operation=operation,
        source=source,
    )


def infer_delimiter(header_line: str, candidates: Sequence[str] = DELIMITER_CANDIDATES) -> str:
    """Pick the delimiter that splits ``header_line`` into the most fields."""
    best = candidates[0]
    best_count = 0
    for candidate in candidates:
        parsed = next(csv.reader([header_line], delimiter=candidate), [])
        if len(parsed) > best_count:
            best = candidate
            best_count = len(parsed)
    return best


def _is_blank(texts: Sequence[str | None]) -> bool:
    """A row is blank when every cell is missing or pure whitespace."""
    return all(text is None or not text.strip() for text in texts)


def _build_row(
    resolution: HeaderResolution,
    schema: IngestSchema,
    texts: Sequence[str | None],
    provenance: RowProvenance,
) -> IngestedRow:
    """Assemble one canonical row; every declared field is always present."""
    values: dict[str, str | None] = dict.fromkeys(schema.field_names)
    extras: dict[str, str | None] = {}
    for index, canonical in enumerate(resolution.columns):
        text = texts[index]
        if canonical is None:
            if schema.extra_columns is ExtraColumnPolicy.RETAIN:
                extras[resolution.normalized[index]] = text
            continue
        values[canonical] = text
    return IngestedRow(provenance=provenance, values=values, extras=extras)


def read_csv(
    path: str | Path,
    schema: IngestSchema,
    *,
    source_label: str | None = None,
    encoding: str | None = None,
    encodings: Iterable[str] = DEFAULT_ENCODINGS,
    delimiter: str | None = None,
) -> IngestResult:
    """Ingest a delimited text file into canonical rows, read-only."""
    operation = "read_csv"
    source_path = Path(path)
    source = safe_source_label(
        source_label if source_label else source_path.name, operation=operation
    )

    chain = (encoding,) if encoding else tuple(encodings)
    text, used_encoding = _decode(source_path, chain, operation=operation, source=source)

    header_line = next((line for line in text.splitlines() if line.strip()), None)
    if header_line is None:
        raise SourceReadError(
            "the source contains no header row", operation=operation, source=source
        )
    used_delimiter = delimiter if delimiter else infer_delimiter(header_line)

    resolution: HeaderResolution | None = None
    header_row = 0
    rows: list[IngestedRow] = []
    skipped_blank = 0

    with source_path.open("r", encoding=used_encoding, newline="") as handle:
        reader = csv.reader(handle, delimiter=used_delimiter)
        previous_line = 0
        for record in reader:
            first_line = previous_line + 1
            previous_line = reader.line_num

            texts = [
                cell_to_text(
                    field,
                    operation=operation,
                    source=source,
                    row=first_line,
                    column=position,
                )
                for position, field in enumerate(record, start=1)
            ]
            if _is_blank(texts):
                # Blank lines ahead of the header are file noise, not data.
                if resolution is not None:
                    skipped_blank += 1
                continue

            if resolution is None:
                resolution = resolve_headers(
                    schema, texts, operation=operation, source=source
                )
                header_row = first_line
                continue

            if len(texts) != resolution.width:
                raise MalformedRowError(
                    f"the row has {len(texts)} field(s) but the header row "
                    f"declares {resolution.width}",
                    operation=operation,
                    source=source,
                    row=first_line,
                )
            rows.append(
                _build_row(
                    resolution,
                    schema,
                    texts,
                    RowProvenance(source=source, source_row=first_line),
                )
            )

    if resolution is None:
        # Reachable when the first non-blank *line* parses to nothing but
        # empty fields, e.g. a row of bare delimiters.
        raise SourceReadError(
            "the source contains no header row", operation=operation, source=source
        )

    return IngestResult(
        source=source,
        source_format=SourceFormat.CSV,
        headers=resolution.headers,
        rows=tuple(rows),
        header_row=header_row,
        skipped_blank_rows=skipped_blank,
        unmapped_headers=resolution.unmapped,
        outcome=IngestOutcome.ROWS if rows else IngestOutcome.HEADER_ONLY,
    )


def _select_worksheet(workbook, worksheet: str | None, *, operation: str, source: str):
    if not workbook.worksheets:
        raise WorksheetSelectionError(
            "the workbook contains no worksheets", operation=operation, source=source
        )
    if worksheet is None:
        return workbook.worksheets[0]
    if worksheet not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise WorksheetSelectionError(
            f"no such worksheet; the workbook offers: {available}",
            operation=operation,
            source=source,
            worksheet=worksheet,
        )
    return workbook[worksheet]


def _row_texts(cells, *, operation: str, source: str, sheet: str, row: int) -> list[str | None]:
    """Convert one worksheet row to canonical text, refusing formula cells."""
    texts: list[str | None] = []
    for position, cell in enumerate(cells, start=1):
        if getattr(cell, "data_type", None) == _FORMULA_DATA_TYPE:
            raise FormulaCellError(
                "the cell holds a formula; ingestion never evaluates formulas "
                "or trusts cached results",
                operation=operation,
                source=source,
                worksheet=sheet,
                row=row,
                column=position,
            )
        texts.append(
            cell_to_text(
                cell.value,
                operation=operation,
                source=source,
                worksheet=sheet,
                row=row,
                column=position,
            )
        )
    return texts


def read_xlsx(
    path: str | Path,
    schema: IngestSchema,
    *,
    source_label: str | None = None,
    worksheet: str | None = None,
) -> IngestResult:
    """Ingest one worksheet into canonical rows, read-only and without saving."""
    operation = "read_xlsx"
    source_path = Path(path)
    source = safe_source_label(
        source_label if source_label else source_path.name, operation=operation
    )

    if not source_path.is_file():
        raise SourceReadError(
            "the source does not exist", operation=operation, source=source
        )
    try:
        workbook = load_workbook(
            source_path,
            read_only=True,
            data_only=False,  # never read a cached formula result as data
            keep_links=False,  # do not retain external workbook links
            keep_vba=False,  # do not retain macros
        )
    except (OSError, BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise SourceReadError(
            f"the workbook could not be opened ({type(exc).__name__})",
            operation=operation,
            source=source,
        ) from None

    try:
        sheet = _select_worksheet(workbook, worksheet, operation=operation, source=source)
        sheet_name = sheet.title

        resolution: HeaderResolution | None = None
        header_row = 0
        rows: list[IngestedRow] = []
        skipped_blank = 0

        for offset, cells in enumerate(sheet.iter_rows(), start=sheet.min_row or 1):
            texts = _row_texts(
                cells, operation=operation, source=source, sheet=sheet_name, row=offset
            )

            if resolution is None:
                trimmed = list(texts)
                while trimmed and trimmed[-1] is None:
                    trimmed.pop()
                if _is_blank(trimmed):
                    continue
                resolution = resolve_headers(
                    schema,
                    trimmed,
                    operation=operation,
                    source=source,
                    worksheet=sheet_name,
                )
                header_row = offset
                continue

            if _is_blank(texts):
                skipped_blank += 1
                continue

            beyond = texts[resolution.width :]
            for extra_offset, value in enumerate(beyond):
                if value is not None:
                    raise MalformedRowError(
                        f"the row carries a value beyond the {resolution.width} "
                        "column(s) the header row declares",
                        operation=operation,
                        source=source,
                        worksheet=sheet_name,
                        row=offset,
                        column=resolution.width + extra_offset + 1,
                    )

            rows.append(
                _build_row(
                    resolution,
                    schema,
                    texts[: resolution.width],
                    RowProvenance(
                        source=source, source_row=offset, worksheet=sheet_name
                    ),
                )
            )

        if resolution is None:
            raise SourceReadError(
                "the worksheet contains no header row",
                operation=operation,
                source=source,
                worksheet=sheet_name,
            )

        return IngestResult(
            source=source,
            source_format=SourceFormat.XLSX,
            headers=resolution.headers,
            rows=tuple(rows),
            header_row=header_row,
            skipped_blank_rows=skipped_blank,
            unmapped_headers=resolution.unmapped,
            worksheet=sheet_name,
            outcome=IngestOutcome.ROWS if rows else IngestOutcome.HEADER_ONLY,
        )
    finally:
        workbook.close()


#: Suffix to reader dispatch, lower-cased.
_READERS = {".csv": SourceFormat.CSV, ".xlsx": SourceFormat.XLSX}


def read_source(path: str | Path, schema: IngestSchema, **options) -> IngestResult:
    """Read a source by file extension, dispatching to the matching reader."""
    source_path = Path(path)
    suffix = source_path.suffix.lower()
    fmt = _READERS.get(suffix)
    if fmt is None:
        supported = ", ".join(sorted(_READERS))
        raise UnsupportedFormatError(
            f"'{suffix or 'no extension'}' is not a supported source format "
            f"(supported: {supported})",
            operation="read_source",
            source=source_path.name,
        )
    if fmt is SourceFormat.CSV:
        return read_csv(source_path, schema, **options)
    return read_xlsx(source_path, schema, **options)


__all__ = [
    "DEFAULT_ENCODINGS",
    "DELIMITER_CANDIDATES",
    "IngestError",
    "infer_delimiter",
    "read_csv",
    "read_source",
    "read_xlsx",
]
